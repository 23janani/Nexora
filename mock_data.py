"""
Mock Data & Demo Dataset Generator
Provides realistic pre-configured candidates spanning high-match, mid-match, underqualified,
duplicate uploads, suspicious repeated content, and AI-styled resumes,
plus realistic Job Descriptions and sample files on disk.
"""

import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from PIL import Image, ImageDraw, ImageFont

SAMPLE_JDS = {
    "senior_ds_genai": {
        "id": "jd-001",
        "title": "Senior Data Scientist - GenAI & MLOps",
        "text": """Job Title: Senior Data Scientist - GenAI & MLOps
Department: Advanced Analytics & AI
Location: San Francisco, CA / Remote

About the Role:
We are seeking an experienced Senior Data Scientist to architect and scale our enterprise generative AI and machine learning infrastructure. You will work closely with cross-functional product and engineering teams to translate high-impact business problems into robust predictive models.

Requirements (Must-Have):
• 4+ years of professional industry experience in data science, predictive modeling, or machine learning engineering.
• Advanced proficiency in Python, including deep familiarity with scientific libraries such as Pandas, NumPy, and Scikit-learn.
• Proven track record of machine learning model development, validation, and evaluation in production environments.
• Experience with SQL and relational databases (PostgreSQL, Snowflake, or BigQuery) for complex data extraction and feature engineering.
• Hands-on experience with cloud infrastructure, specifically AWS (EC2, S3, SageMaker, or Lambda) for model training and deployment.
• Solid foundation in statistics, probability, and exploratory data analysis.

Preferred Qualifications (Nice-to-Have):
• Experience with Deep Learning frameworks (PyTorch or TensorFlow) and LLM orchestration (LangChain, Hugging Face).
• Working knowledge of MLOps best practices, Docker containerization, Kubernetes, and automated CI/CD pipelines.
• Master's or Ph.D. degree in Computer Science, Statistics, Mathematics, Data Science, or related quantitative field.
• Excellent communication skills with ability to present technical insights to non-technical leadership.
"""
    },
    "ml_infra_eng": {
        "id": "jd-002",
        "title": "Machine Learning Platform Engineer",
        "text": """Job Title: Machine Learning Platform Engineer
Department: Core ML Infrastructure
Location: Seattle, WA

Role Overview:
Build distributed platforms supporting model serving, real-time feature stores, and automated training pipelines.

Requirements (Must-Have):
• 3+ years engineering scalable cloud pipelines on AWS or GCP.
• Deep expertise with Docker containerization and Kubernetes cluster management.
• High proficiency in Python and Go or C++.
• Demonstrated experience with MLOps frameworks like Kubeflow or MLflow.
• Strong foundation in SQL and distributed data stores.

Preferred Qualifications (Nice-to-Have):
• Experience optimizing deep learning inference engines (TensorRT, ONNX).
• BS/MS in Computer Science.
"""
    },
    "data_analyst": {
        "id": "jd-003",
        "title": "Business Intelligence & Data Analyst",
        "text": """Job Title: Business Intelligence & Data Analyst
Department: Business Operations
Location: New York, NY

Requirements (Must-Have):
• 2+ years experience building executive BI dashboards in Tableau or Power BI.
• Advanced SQL skills (window functions, subqueries, CTEs, performance tuning).
• Python data manipulation skills using Pandas and NumPy.
• Demonstrated experience driving business KPIs and reporting.

Preferred Qualifications:
• Bachelor's degree in Economics, Business Analytics, or Statistics.
• Familiarity with statistical hypothesis testing and A/B experiments.
"""
    }
}

SAMPLE_RESUMES = {
    "elena_vance": {
        "id": "cand-001",
        "name": "Dr. Elena Vance",
        "target_role": "Senior Data Scientist",
        "expected_score_range": "88-94%",
        "text": """Dr. Elena Vance, Ph.D.
Email: elena.vance@alum.stanford.edu | Phone: (415) 555-0194 | Location: San Francisco, CA
LinkedIn: linkedin.com/in/elenavance-ds | GitHub: github.com/elenavance-ai

PROFESSIONAL SUMMARY
Senior Data Scientist with 6+ years of specialized experience architecting predictive machine learning pipelines, LLM solutions, and cloud-native MLOps architectures. Proven track record deploying models serving 5M+ daily requests on AWS with high reliability and low latency.

EDUCATION
Ph.D. in Computer Science (Machine Learning Specialization)
Stanford University, Stanford, CA | 2018
• Dissertation: Efficient Transformer Architectures for Contextual Language Modeling
• Master of Science in Statistics, Stanford University | 2015

WORK EXPERIENCE
Staff Data Scientist | Apex Machine Intelligence, San Francisco, CA
2021 - Present
• Spearheaded predictive models using scikit-learn, XGBoost, and PyTorch for customer risk scoring, lifting prediction accuracy by 24%.
• Architected end-to-end cloud deployment pipelines utilizing AWS EC2, S3, and SageMaker with Docker containerization.
• Designed automated MLOps CI/CD pipelines incorporating MLflow and Kubernetes, reducing retraining turnaround time from 2 weeks to 4 hours.
• Wrote high-performance PostgreSQL and Snowflake queries for real-time feature store ingestion, handling 100GB+ daily transaction streams.

Senior Machine Learning Engineer | Stratum Data Labs, San Jose, CA
2018 - 2021
• Built predictive models using scikit-learn and LightGBM for e-commerce conversion forecasting.
• Maintained strong Python-based data analysis workflows with Pandas, NumPy, and SciPy for exploratory data modeling and hypothesis testing.
• Deployed microservices on AWS Lambda and FastAPI handling real-time inference with sub-40ms response latency.
• Mentored 5 junior data scientists and established team coding standards and testing protocols.

TECHNICAL SKILLS
• Programming Languages: Python, SQL, R, Bash
• Frameworks & Libraries: PyTorch, TensorFlow, Scikit-learn, Pandas, NumPy, SciPy, Hugging Face, XGBoost, FastAPI
• Cloud & MLOps: AWS (EC2, S3, SageMaker, Lambda), Docker, Kubernetes, MLflow, CI/CD, Git
• Databases: PostgreSQL, Snowflake, BigQuery, Redis
• Concepts: Supervised Learning, Deep Learning, Natural Language Processing, MLOps, A/B Testing, Feature Engineering

KEY PROJECTS
• Enterprise LLM Knowledge Retrieval System: Implemented vector search embeddings and contextual retrieval pipelines reducing hallucination rates by 35%.
• Real-time Churn Forecasting Engine: Built multi-task neural network on AWS SageMaker, decreasing customer churn by $3.2M annually.

CERTIFICATIONS
• AWS Certified Machine Learning - Specialty
• Databricks Certified Machine Learning Professional
"""
    },
    "marcus_chen": {
        "id": "cand-002",
        "name": "Marcus Chen",
        "target_role": "Data Scientist / ML Engineer",
        "expected_score_range": "78-85%",
        "text": """Marcus Chen
Email: marcus.chen.dev@gmail.com | Phone: (206) 555-7382 | Location: Seattle, WA

SUMMARY
Results-driven Data Scientist with 4 years of industry experience developing statistical models, data analytics pipelines, and machine learning solutions using Python, SQL, and AWS.

EDUCATION
Master of Science in Data Science
University of Washington, Seattle, WA | 2020
Bachelor of Science in Mathematics
University of Washington, Seattle, WA | 2018

WORK EXPERIENCE
Data Scientist | CloudScale Analytics, Seattle, WA
2021 - Present
• Built predictive models using scikit-learn and random forests to forecast supply chain inventory demands across 50 distribution centers.
• Conducted extensive Python data analysis using Pandas, NumPy, and Matplotlib to extract operational insights for executive leadership.
• Collaborated with engineering teams to deploy containerized models using Docker on AWS EC2 instances.
• Developed optimized SQL queries in PostgreSQL and Snowflake to extract feature tables from multi-terabyte data warehouses.

Associate Data Analyst | BluePeak Solutions, Bellevue, WA
2020 - 2021
• Created interactive Tableau dashboards tracking customer engagement metrics and monthly recurring revenue.
• Automated weekly data hygiene pipelines using Python scripts and SQL views, reducing manual reporting hours by 70%.

TECHNICAL SKILLS
• Languages: Python, SQL
• Libraries: Scikit-learn, Pandas, NumPy, Matplotlib, Seaborn, Statsmodels
• Cloud & Tools: AWS (EC2, S3), Docker, Git, Tableau, Jupyter
• Databases: PostgreSQL, MySQL, Snowflake

PROJECTS
• Predictive Lead Scoring: Developed logistic regression and gradient boosting models in Python achieving 0.86 AUC-ROC.
• NLP Sentiment Classifier: Fine-tuned DistilBERT model for customer ticket routing using Hugging Face and PyTorch.
"""
    },
    "aisha_patel": {
        "id": "cand-003",
        "name": "Aisha Patel",
        "target_role": "Data Analyst (Transitioning to DS)",
        "expected_score_range": "58-65%",
        "text": """Aisha Patel
Email: aisha.patel@outlook.com | Phone: (512) 555-9012 | Location: Austin, TX

PROFESSIONAL SUMMARY
Analytical and detail-oriented Data Analyst with 3 years of experience specializing in complex SQL querying, dashboard development in Tableau, and exploratory Python data analysis. Seeking to transition into Machine Learning and Data Science.

EDUCATION
Bachelor of Science in Economics & Minor in Statistics
University of Texas at Austin | 2021

EXPERIENCE
Data Analyst | LoneStar Financial Services, Austin, TX
2021 - Present
• Authored advanced SQL queries involving multi-table joins, subqueries, and window functions on PostgreSQL database.
• Performed exploratory data analysis in Python using Pandas and NumPy to identify customer attrition trends.
• Designed and published 15+ interactive executive dashboards in Tableau.
• Ran hypothesis testing and statistical power calculations for marketing campaigns.

Business Intelligence Intern | FinMetrics Austin, TX
2020 - 2021
• Assisted in cleaning and validating transaction data using Excel and Python.
• Maintained KPI reporting tracking quarterly revenue and conversion rates.

SKILLS
• Languages: SQL, Python (Pandas, NumPy, Matplotlib)
• BI Tools: Tableau, Power BI, Excel
• Databases: PostgreSQL, MySQL
• Statistical Methods: Regression, Hypothesis Testing, ANOVA, Descriptive Statistics
"""
    },
    "jordan_miller": {
        "id": "cand-004",
        "name": "Jordan Miller",
        "target_role": "Frontend / Web Developer",
        "expected_score_range": "30-40%",
        "text": """Jordan Miller
Email: jordan.miller@techworks.io | Phone: (312) 555-4819 | Location: Chicago, IL

PROFILE
Frontend Web Developer with 3 years of experience building modern, responsive web applications using React, JavaScript, HTML5, CSS3, and Node.js.

EDUCATION
Bachelor of Arts in Graphic Design & Web Communications
University of Illinois Chicago | 2021

PROFESSIONAL EXPERIENCE
Frontend Developer | WebPulse Studios, Chicago, IL
2021 - Present
• Developed single-page applications using React, TypeScript, and Tailwind CSS.
• Integrated REST APIs with Redux for state management.
• Optimized web page performance, achieving 95+ Google Lighthouse scores.
• Collaborated with UX designers in Figma to build accessible, mobile-first design systems.

Junior Web Developer | Creativo Digital, Chicago, IL
2020 - 2021
• Built landing pages with HTML5, CSS3, and vanilla JavaScript.
• Maintained client websites using WordPress and PHP.

SKILLS
• Frontend: JavaScript, TypeScript, React, HTML5, CSS3, Sass, Tailwind CSS
• Backend/Tools: Node.js, Express, Git, Figma, Webpack, Vite
"""
    },
    "david_ross_ai": {
        "id": "cand-005",
        "name": "David Ross",
        "target_role": "Senior Data Scientist (AI-Style Resume)",
        "expected_score_range": "75-84%",
        "text": """David Ross
Email: david.ross.ai.expert@techmail.com | Phone: (617) 555-8822 | Location: Boston, MA

SUMMARY
A results-driven data scientist showcasing a proven track record of driving impactful business outcomes. Spearheaded the development of robust and scalable machine learning pipelines that seamlessly integrated with cutting-edge cloud infrastructure. A testament to technical excellence and cross-functional leadership.

EDUCATION
Master of Science in Data Science
Boston University | 2020

WORK EXPERIENCE
Senior Data Scientist | Pinnacle Dynamics, Boston, MA
2021 - Present
• Spearheaded the development of predictive models leveraging state-of-the-art scikit-learn architectures.
• Played a pivotal role in facilitating cross-functional collaboration across executive stakeholders.
• Meticulously crafted high-throughput data processing workflows with Python, Pandas, and NumPy.
• Seamlessly integrated machine learning inference modules into AWS cloud environments.
• A testament to innovation, delivering a holistic approach toward predictive analytics.

Data Scientist | OmniTech Labs, Cambridge, MA
2020 - 2021
• Leveraged cutting-edge methodologies to train supervised regression algorithms.
• Fostered cross-functional collaboration to optimize transactional databases in PostgreSQL.
• Orchestrated predictive modeling pipelines driving impactful business outcomes.

TECHNICAL SKILLS
Python, Scikit-learn, Pandas, NumPy, AWS, SQL, Machine Learning, Data Modeling
"""
    },
    "elena_vance_duplicate": {
        "id": "cand-006",
        "name": "Elena Vance (Duplicate Upload)",
        "target_role": "Senior Data Scientist",
        "expected_score_range": "Duplicate Flag (98% Sim)",
        "text": """Dr. Elena Vance, Ph.D.
Email: elena.vance@alum.stanford.edu | Phone: (415) 555-0194 | Location: San Francisco, CA

PROFESSIONAL SUMMARY
Senior Data Scientist with 6+ years of specialized experience architecting predictive machine learning pipelines, LLM solutions, and cloud-native MLOps architectures. Proven track record deploying models serving 5M+ daily requests on AWS with high reliability and low latency.

EDUCATION
Ph.D. in Computer Science (Machine Learning Specialization), Stanford University, 2018
Master of Science in Statistics, Stanford University, 2015

WORK EXPERIENCE
Staff Data Scientist | Apex Machine Intelligence, San Francisco, CA (2021 - Present)
• Spearheaded predictive models using scikit-learn, XGBoost, and PyTorch for customer risk scoring, lifting prediction accuracy by 24%.
• Architected end-to-end cloud deployment pipelines utilizing AWS EC2, S3, and SageMaker with Docker containerization.
• Designed automated MLOps CI/CD pipelines incorporating MLflow and Kubernetes.
• Wrote high-performance PostgreSQL and Snowflake queries for real-time feature store ingestion.

TECHNICAL SKILLS
Python, SQL, PyTorch, Scikit-learn, Pandas, NumPy, AWS, Docker, Kubernetes, PostgreSQL
"""
    },
    "spam_repeated": {
        "id": "cand-007",
        "name": "Spam / Repeated Content Applicant",
        "target_role": "Applicant with Keyword Stuffing",
        "expected_score_range": "Content Flagged",
        "text": """Alex Mercer
Email: alex.mercer@inbox.net | Phone: (408) 555-3311

SUMMARY
Machine learning machine learning python python data science data science machine learning.
I have extensive experience with machine learning machine learning python python sql sql.

EXPERIENCE
EXPERIENCE
EXPERIENCE
• Developed machine learning models using python and machine learning algorithms.
• Developed machine learning models using python and machine learning algorithms.
• Developed machine learning models using python and machine learning algorithms.
• Machine learning models were developed using python data science libraries and machine learning.
• Machine learning models were developed using python data science libraries and machine learning.

SKILLS
Python, Python, Python, Machine Learning, Machine Learning, Machine Learning, SQL, SQL, SQL, Pandas, Pandas, Pandas.
"""
    }
}


def create_sample_files(target_dir: str = "./sample_files"):
    """
    Generate actual files on disk for user testing:
    - PDF resumes (via reportlab)
    - DOCX resumes (via python-docx)
    - TXT resumes
    - PNG resume (image format for OCR testing via Pillow)
    - Sample JD text file
    """
    os.makedirs(target_dir, exist_ok=True)

    # 1. Sample JD File
    jd_path = os.path.join(target_dir, "Job_Description_Senior_Data_Scientist.txt")
    with open(jd_path, "w", encoding="utf-8") as f:
        f.write(SAMPLE_JDS["senior_ds_genai"]["text"])

    # 2. Sample Resumes in TXT
    import re
    for key, data in SAMPLE_RESUMES.items():
        safe_name = re.sub(r'[^a-zA-Z0-9_]', '_', data['name'])
        fname = f"Resume_{safe_name}.txt"
        fpath = os.path.join(target_dir, fname)
        with open(fpath, "w", encoding="utf-8") as f:
            f.write(data["text"])

    # 3. Sample DOCX resume
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Dr. Elena Vance, Ph.D.", level=0)
        doc.add_paragraph("Email: elena.vance@alum.stanford.edu | Location: San Francisco, CA")
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph("Senior Data Scientist with 6+ years experience architecting ML models on AWS.")
        doc.add_heading("Work Experience", level=1)
        doc.add_paragraph("Staff Data Scientist at Apex Machine Intelligence (2021 - Present)")
        doc.add_paragraph("• Built predictive models using scikit-learn and PyTorch on AWS EC2 and SageMaker.")
        doc.add_paragraph("• Deployed Docker containers with Kubernetes CI/CD pipelines.")
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("Python, PyTorch, Scikit-learn, AWS, Docker, Kubernetes, PostgreSQL, SQL")
        doc.save(os.path.join(target_dir, "Resume_Elena_Vance.docx"))
    except Exception as e:
        print(f"Could not create sample docx: {e}")

    # 4. Sample PDF resume
    try:
        pdf_path = os.path.join(target_dir, "Resume_Elena_Vance.pdf")
        doc_pdf = SimpleDocTemplate(pdf_path, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("TitleStyle", parent=styles["Heading1"], fontSize=18, leading=22, textColor="#1e293b")
        h2_style = ParagraphStyle("H2Style", parent=styles["Heading2"], fontSize=13, leading=16, textColor="#0f172a", spaceBefore=10, spaceAfter=4)
        body_style = ParagraphStyle("BodyStyle", parent=styles["Normal"], fontSize=9.5, leading=13, textColor="#334155")

        story = [
            Paragraph("Dr. Elena Vance, Ph.D. - Senior Data Scientist", title_style),
            Paragraph("Email: elena.vance@alum.stanford.edu | Phone: (415) 555-0194 | Location: San Francisco, CA", body_style),
            Spacer(1, 10),
            Paragraph("Professional Summary", h2_style),
            Paragraph("Senior Data Scientist with 6+ years of specialized experience architecting predictive machine learning pipelines, LLM solutions, and cloud-native MLOps architectures. Proven track record deploying models serving 5M+ daily requests on AWS with high reliability and low latency.", body_style),
            Spacer(1, 8),
            Paragraph("Education", h2_style),
            Paragraph("Ph.D. in Computer Science (Machine Learning Specialization), Stanford University, 2018.<br/>Master of Science in Statistics, Stanford University, 2015.", body_style),
            Spacer(1, 8),
            Paragraph("Work Experience", h2_style),
            Paragraph("<b>Staff Data Scientist</b> - Apex Machine Intelligence (2021 - Present)<br/>• Spearheaded predictive models using scikit-learn, XGBoost, and PyTorch for customer risk scoring.<br/>• Architected cloud deployment pipelines utilizing AWS EC2, S3, and SageMaker with Docker and Kubernetes.<br/>• Wrote high-performance PostgreSQL and Snowflake queries for real-time feature store ingestion.", body_style),
            Spacer(1, 8),
            Paragraph("Skills", h2_style),
            Paragraph("Python, PyTorch, Scikit-learn, Pandas, NumPy, AWS, Docker, Kubernetes, SQL, PostgreSQL, Snowflake, MLOps.", body_style)
        ]
        doc_pdf.build(story)
    except Exception as e:
        print(f"Could not create sample pdf: {e}")

    # 5. Sample PNG Image Resume (for OCR testing)
    try:
        img_path = os.path.join(target_dir, "Resume_Marcus_Chen.png")
        img = Image.new("RGB", (1000, 1100), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        lines = [
            "Marcus Chen - Data Scientist & ML Engineer",
            "Email: marcus.chen.dev@gmail.com | Phone: (206) 555-7382 | Seattle, WA",
            "-------------------------------------------------------------------------",
            "SUMMARY:",
            "Results-driven Data Scientist with 4 years of industry experience developing",
            "statistical models, data analytics pipelines, and ML solutions using Python, SQL, AWS.",
            "",
            "EDUCATION:",
            "Master of Science in Data Science, University of Washington (2020)",
            "",
            "EXPERIENCE:",
            "Data Scientist at CloudScale Analytics (2021 - Present)",
            "- Built predictive models using scikit-learn and random forests.",
            "- Conducted extensive Python data analysis using Pandas, NumPy, and Matplotlib.",
            "- Deployed containerized models using Docker on AWS EC2 instances.",
            "- Developed optimized SQL queries in PostgreSQL and Snowflake.",
            "",
            "SKILLS:",
            "Python, SQL, Scikit-learn, Pandas, NumPy, AWS (EC2, S3), Docker, Git, Tableau."
        ]
        y = 40
        for line in lines:
            draw.text((50, y), line, fill=(20, 20, 20))
            y += 35
        img.save(img_path)
    except Exception as e:
        print(f"Could not create sample png: {e}")

    print(f"All sample files (PDF, DOCX, TXT, PNG) ready in {target_dir}")
